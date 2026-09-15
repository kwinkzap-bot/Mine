"""Append a month of sessions to the CPR sheet, analysis columns read off the chart.

The hand-analysed January rows of "Bactest Nifty CPR.xlsx" were written by
eye. Every month since is written the same way the February rows were:
this script adds one row per session with the chart's reading of the
analysis columns (B-G, the same reading the Trend page shows under the
sheet's word), a note in O saying so, and the sheet's P&L formula in M —
then scripts/propose_cpr_trades.py fills the trade columns by rule, and
scripts/import_cpr_manual.py takes the sheet into the Trend page.

    python scripts/add_cpr_sessions.py ~/Downloads/"Bactest Nifty CPR.xlsx" \\
        --sheet 2026 --from 2026-03-01 --to 2026-03-31
    python scripts/propose_cpr_trades.py ~/Downloads/"Bactest Nifty CPR.xlsx" \\
        --sheet 2026 --from 2026-03-01 --to 2026-03-31
    python scripts/import_cpr_manual.py ~/Downloads/"Bactest Nifty CPR.xlsx" --sheet 2026

A session already on the sheet is left alone (its date is skipped), so the
range may overlap what is there. Rows go after the last dated row, in date
order; the sessions are whichever days Fyers has 5-minute bars for.
"""

import argparse
import os
import sys
from datetime import date, datetime

import openpyxl
from openpyxl.styles import PatternFill

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src'))

from trading_app.service import cpr_backtest_service as svc  # noqa: E402
from propose_cpr_trades import bars_for  # noqa: E402

FILL = PatternFill('solid', fgColor='E2EFDA')   # light green: analysis written from the chart
NOTE = 'Analysis added from the chart (Fyers 5-min)'

# The chart's level names in the sheet's words.
_LEVEL_WORDS = {'PDH': 'Prev High', 'PDL': 'Prev Low', 'Cam R3': 'R3(cam)', 'Cam S3': 'S3(cam)'}


def candle_words(fc):
    """The chart's first-candle reading in the sheet's vocabulary —
    'Strong small candle red near Prev Low', 'In decision candle(doji) near CPR'."""
    text = fc['text']
    level = None
    if ' · touched ' in text:
        level = text.split(' · touched ')[1].split(', ')[0]
    elif ' · near ' in text:
        level = text.split(' · near ')[1]
    head = text.split(' · ')[0]
    if fc['colour'] == 'Doji':
        words = 'In decision candle(doji)'
    else:
        parts = head.split()               # e.g. ['Strong', 'small', 'Red']
        colour = parts[-1].lower()
        adj = ' '.join(p.lower() for p in parts[:-1])
        words = (adj.capitalize() + ' ' if adj else '') + f'candle {colour}'
        words = words[0].upper() + words[1:]
    if level:
        words += ' near ' + _LEVEL_WORDS.get(level, level)
    return words


def pnl_formula(rn):
    return (f'=IF(OR(I{rn}="", J{rn}="", L{rn}=""), "", IF(L{rn}="Target", IF(J{rn}>I{rn}, J{rn}-I{rn}, I{rn}-J{rn}), '
            f'IF(AND(L{rn}="SL", K{rn}<>""), IF(J{rn}>I{rn}, K{rn}-I{rn}, I{rn}-K{rn}), "")))')


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('xlsx')
    ap.add_argument('--sheet', default=None)
    ap.add_argument('--symbol', default='NIFTY')
    ap.add_argument('--from', dest='first', required=True)
    ap.add_argument('--to', dest='last', required=True)
    ap.add_argument('--dry-run', action='store_true', help='print, do not write')
    args = ap.parse_args(argv)
    first, last = date.fromisoformat(args.first), date.fromisoformat(args.last)
    last = min(last, date.today())

    wb = openpyxl.load_workbook(args.xlsx)
    ws = wb[args.sheet] if args.sheet else wb.worksheets[0]
    have, last_row = set(), 1
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        d = row[0].value
        if isinstance(d, datetime):
            have.add(d.date().isoformat())
            last_row = row[0].row

    daily, intraday = bars_for(args.symbol, first, last)
    sessions = sorted(ds for ds in intraday
                      if first <= date.fromisoformat(ds) <= last and ds not in have)
    if not sessions:
        print('nothing to do: every session in range is already on the sheet')
        return
    res = svc.analyse([{'date': ds} for ds in sessions], daily, intraday)

    rn, written = last_row, 0
    for r in res['rows']:
        c = r['chart']
        if not c:
            print(f"{r['date']}: {r.get('error')} — skipped")
            continue
        fc = c['first_candle']
        vals = {
            2: f"Price {c['price_vs_daily']}",
            3: f"Price {c['price_vs_hourly']}" if c['price_vs_hourly'] else None,
            4: c['cpr_type'], 5: c['cpr_direction'], 6: candle_words(fc), 7: c['boxes'],
        }
        note = (f"{NOTE}: CPR {abs(c['levels']['tc'] - c['levels']['bc']):.1f} pts ({c['width_pct']}%), "
                f"09:15 candle O {fc['open']} H {fc['high']} L {fc['low']} C {fc['close']}")
        print(r['date'], [vals[k] for k in range(2, 8)])
        if args.dry_run:
            continue
        rn += 1
        a = ws.cell(rn, 1)
        a.value = datetime.strptime(r['date'], '%Y-%m-%d')
        a.number_format = 'dd/MM/yyyy'
        a.fill = FILL
        for col, v in vals.items():
            cell = ws.cell(rn, col)
            cell.value = v
            cell.fill = FILL
        ws.cell(rn, 13).value = pnl_formula(rn)
        ws.cell(rn, 15).value = note
        written += 1
    if not args.dry_run:
        wb.save(args.xlsx)
        print(f'wrote {written} sessions to {os.path.basename(args.xlsx)}')


if __name__ == '__main__':
    main()
